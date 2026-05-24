from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF

from app.ingestion.parsers import DOCXParser, MarkdownParser, PDFParser, TXTParser


def _make_pdf(tmp_path: Path, pages: list[str]) -> str:
    """Generate a multi-page PDF with given page texts using fpdf2."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.set_xy(10, 10)
        pdf.cell(0, 10, text)
    path = tmp_path / "test.pdf"
    pdf.output(str(path))
    return str(path)


def _make_docx(tmp_path: Path, paragraphs: list[str]) -> str:
    """Generate a DOCX with given paragraphs using python-docx."""
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "test.docx"
    doc.save(str(path))
    return str(path)


def test_txt_parser_reads_content(tmp_path: Path) -> None:
    file_path = tmp_path / "test.txt"
    file_path.write_text("hello world", encoding="utf-8")

    result = TXTParser().parse("doc_1", "public", str(file_path), title="Test")

    assert result.document_id == "doc_1"
    assert result.tenant_id == "public"
    assert result.title == "Test"
    assert result.source_type == "txt"
    assert result.extracted_text == "hello world"
    assert len(result.sections) == 1
    assert result.sections[0].content == "hello world"


def test_markdown_parser_produces_sections(tmp_path: Path) -> None:
    md_content = """# Title

Some intro text.

## Section 1

Content for section one.

## Section 2

Content for section two.
"""
    file_path = tmp_path / "test.md"
    file_path.write_text(md_content, encoding="utf-8")

    result = MarkdownParser().parse("doc_2", "tenant-a", str(file_path), title="Doc")

    assert result.source_type == "md"
    assert result.document_id == "doc_2"
    assert len(result.sections) >= 2


def test_markdown_parser_handles_plain_text_without_headings(tmp_path: Path) -> None:
    file_path = tmp_path / "noheadings.md"
    file_path.write_text("Just a plain text without any markdown headings.", encoding="utf-8")

    result = MarkdownParser().parse("doc_3", "public", str(file_path))

    assert len(result.sections) == 1
    assert "headings" in result.sections[0].content


def test_markdown_parser_inline_without_heading_sets_section_heading(tmp_path: Path) -> None:
    """First inline token becomes heading, subsequent inline tokens become content."""
    file_path = tmp_path / "inline_first.md"
    file_path.write_text("First paragraph before any heading.\n\nSecond line.", encoding="utf-8")

    result = MarkdownParser().parse("doc_4", "public", str(file_path))

    assert len(result.sections) >= 1
    assert result.sections[0].heading == "First paragraph before any heading."
    assert result.sections[0].content == "Second line."


def test_pdf_parser_reads_single_page(tmp_path: Path) -> None:
    source = _make_pdf(tmp_path, ["Hello PDF World"])
    result = PDFParser().parse("doc_p1", "public", source, title="PDF Test")

    assert result.document_id == "doc_p1"
    assert result.source_type == "pdf"
    assert result.title == "PDF Test"
    assert result.extracted_text == "Hello PDF World"
    assert len(result.sections) == 1
    assert result.sections[0].page_number == 1
    assert result.sections[0].heading == "Page 1"


def test_pdf_parser_multi_page(tmp_path: Path) -> None:
    source = _make_pdf(tmp_path, ["Page one content", "Page two content"])
    result = PDFParser().parse("doc_p2", "tenant-a", source)

    assert result.document_id == "doc_p2"
    assert len(result.sections) == 2
    assert result.sections[0].page_number == 1
    assert result.sections[1].page_number == 2
    assert result.sections[0].content == "Page one content"
    assert result.sections[1].content == "Page two content"
    assert result.metadata == {"page_count": 2}


def test_pdf_parser_produces_extracted_text(tmp_path: Path) -> None:
    source = _make_pdf(tmp_path, ["Page one", "Page two"])
    result = PDFParser().parse("doc_p4", "public", source)
    assert result.extracted_text == "Page one\n\nPage two"


def test_pdf_parser_skips_empty_pages(tmp_path: Path) -> None:
    source = _make_pdf(tmp_path, ["Has content", "", "Also content"])
    result = PDFParser().parse("doc_p3", "public", source)

    assert len(result.sections) == 2  # empty page 2 is skipped
    assert result.sections[0].page_number == 1
    assert result.sections[1].page_number == 3


def test_docx_parser_reads_paragraphs(tmp_path: Path) -> None:
    source = _make_docx(tmp_path, ["First paragraph", "Second paragraph"])
    result = DOCXParser().parse("doc_d1", "public", source, title="DOCX Test")

    assert result.document_id == "doc_d1"
    assert result.source_type == "docx"
    assert result.title == "DOCX Test"
    assert result.metadata == {"paragraph_count": 2}
    assert "First paragraph" in result.extracted_text
    assert "Second paragraph" in result.extracted_text
    assert len(result.sections) == 1


def test_docx_parser_skips_empty_paragraphs(tmp_path: Path) -> None:
    source = _make_docx(tmp_path, ["First", "", "Third"])
    result = DOCXParser().parse("doc_d2", "public", source)

    assert result.metadata == {"paragraph_count": 2}  # empty paragraph excluded
    assert "First" in result.extracted_text
    assert "Third" in result.extracted_text
